#!/usr/bin/env python3
import http.server
import socketserver
import json
import os
import sys
import io
import re
import tempfile
import shutil
import subprocess
import base64
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from lxml import etree
import latex2mathml.converter
try:
    from geometry_engine import DIAGRAM_PRESETS, render_diagram_png
except ImportError:
    from live_editor.geometry_engine import DIAGRAM_PRESETS, render_diagram_png

PORT = 8080
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(BASE_DIR)

# Load Microsoft's official MathML to OMML XSLT
XSL_PATH = '/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl'
if not os.path.exists(XSL_PATH):
    print(f"Warning: XSL not found at {XSL_PATH}")
    xslt_transform = None
else:
    xslt_doc = etree.parse(XSL_PATH)
    xslt_transform = etree.XSLT(xslt_doc)

def latex_to_omml_xml(latex_str):
    mathml = latex2mathml.converter.convert(latex_str)
    if xslt_transform:
        mml_root = etree.fromstring(mathml.encode('utf-8'))
        omml = xslt_transform(mml_root)
        return etree.tostring(omml.getroot(), encoding='utf-8').decode('utf-8')
    return f"<m:oMath xmlns:m=\"http://schemas.openxmlformats.org/officeDocument/2006/math\"><m:r><m:t>{latex_str}</m:t></m:r></m:oMath>"

def set_run_font(run, font_name="Khmer OS Battambang", font_size=11, color=None, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

def add_mixed_text(p, text_with_math, default_font="Khmer OS Battambang", font_size=10.5, text_color=RGBColor(30, 41, 59), bold=False):
    pattern = re.compile(r'(\$\$.*?\$\$|\$.*?\$)')
    parts = pattern.split(text_with_math)
    for part in parts:
        if not part:
            continue
        if part.startswith('$$') and part.endswith('$$'):
            latex_code = part[2:-2].strip()
            try:
                omml_xml = latex_to_omml_xml(latex_code)
                p._p.append(parse_xml(omml_xml))
            except Exception:
                r = p.add_run(latex_code)
                set_run_font(r, default_font, font_size, text_color, bold=bold)
        elif part.startswith('$') and part.endswith('$'):
            latex_code = part[1:-1].strip()
            try:
                omml_xml = latex_to_omml_xml(latex_code)
                p._p.append(parse_xml(omml_xml))
            except Exception:
                r = p.add_run(latex_code)
                set_run_font(r, default_font, font_size, text_color, bold=bold)
        else:
            r = p.add_run(part)
            set_run_font(r, default_font, font_size, text_color, bold=bold)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_val = edge_data.get('val', 'single')
            b_sz = edge_data.get('sz', '4')
            b_space = edge_data.get('space', '0')
            b_color = edge_data.get('color', 'auto')
            border_el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{b_val}" w:sz="{b_sz}" w:space="{b_space}" w:color="{b_color}"/>')
            tcBorders.append(border_el)
        else:
            border_el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(border_el)
    tcPr.append(tcBorders)

def add_badge(p, text):
    run = p.add_run(f" {text} ")
    set_run_font(run, "Khmer OS Battambang", 9, RGBColor(255, 255, 255), bold=True)
    rPr = run._r.get_or_add_rPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="C3233C"/>')
    rPr.append(shd)
    sp = p.add_run(" ")
    set_run_font(sp, "Khmer OS Battambang", 10.5)

def build_docx_from_data(data):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)

    # 1. Header Banner
    t_head = doc.add_table(rows=1, cols=3)
    t_head.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_head.autofit = False
    c0, c1, c2 = t_head.rows[0].cells
    c0.width = Cm(3.6)
    c1.width = Cm(9.4)
    c2.width = Cm(5.0)

    set_cell_shading(c0, "00969B")
    set_cell_margins(c0, 120, 120, 150, 150)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_num = p0.add_run(data.get("lesson_num", "មេរៀនទី ..."))
    set_run_font(r_num, "Khmer OS Muol Light", 11, RGBColor(255, 255, 255), bold=True)

    set_cell_margins(c1, 100, 100, 150, 100)
    p1 = c1.paragraphs[0]
    r_title = p1.add_run(data.get("lesson_title", "ចំណងជើងមេរៀន"))
    set_run_font(r_title, "Khmer OS Muol Light", 17, RGBColor(0, 105, 110), bold=True)

    set_cell_margins(c2, 140, 100, 50, 50)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sub = p2.add_run(data.get("subtitle", "♣ ក្លឹបសិក្សាគណិតវិទ្យា ភី ស៊ី​ អឹម"))
    set_run_font(r_sub, "Khmer OS Battambang", 8.5, RGBColor(115, 125, 135))

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(3)
    sp.paragraph_format.space_after = Pt(3)

    # 2. Section Banner
    t_sec = doc.add_table(rows=1, cols=3)
    t_sec.alignment = WD_TABLE_ALIGNMENT.LEFT
    t_sec.autofit = False
    sc0, sc1, sc2 = t_sec.rows[0].cells
    sc0.width = Cm(1.1)
    sc1.width = Cm(0.15)
    sc2.width = Cm(16.75)

    set_cell_shading(sc0, "0F2C59")
    set_cell_margins(sc0, 70, 70, 60, 60)
    sp0 = sc0.paragraphs[0]
    sp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_snum = sp0.add_run(data.get("section_num", "១"))
    set_run_font(r_snum, "Khmer OS Muol Light", 12, RGBColor(255, 255, 255), bold=True)

    set_cell_shading(sc1, "D4A373")
    set_cell_margins(sc1, 0, 0, 0, 0)

    set_cell_margins(sc2, 70, 70, 150, 100)
    sp2 = sc2.paragraphs[0]
    r_stitle = sp2.add_run(data.get("section_title", "ចំណងជើងផ្នែកធំ"))
    set_run_font(r_stitle, "Khmer OS Muol Light", 13, RGBColor(195, 35, 60), bold=True)

    sp2_p = doc.add_paragraph()
    sp2_p.paragraph_format.space_before = Pt(2)
    sp2_p.paragraph_format.space_after = Pt(2)

    # Subsection
    subsec_title = data.get("subsec_title")
    if subsec_title:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(2)
        p_sub.paragraph_format.space_after = Pt(4)
        r_bar = p_sub.add_run("▌ ")
        set_run_font(r_bar, "Khmer OS Muol Light", 11, RGBColor(15, 44, 89))
        r_txt = p_sub.add_run(subsec_title)
        set_run_font(r_txt, "Khmer OS Muol Light", 11, RGBColor(15, 44, 89), bold=True)

    # 3. Theory Box
    theories = data.get("theories", [])
    if theories:
        t_box = doc.add_table(rows=1, cols=1)
        t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_box.autofit = False
        bcell = t_box.rows[0].cells[0]
        bcell.width = Cm(18.0)
        set_cell_shading(bcell, "F7F4EE")
        set_cell_margins(bcell, 140, 140, 200, 200)
        set_cell_border(bcell, left={'val': 'single', 'sz': '24', 'color': '0F2C59'},
                               top={'val': 'single', 'sz': '4', 'color': 'E0D9CE'},
                               right={'val': 'single', 'sz': '4', 'color': 'E0D9CE'},
                               bottom={'val': 'single', 'sz': '4', 'color': 'E0D9CE'})

        p_bhead = bcell.paragraphs[0]
        p_bhead.paragraph_format.space_after = Pt(4)
        r_bh = p_bhead.add_run("★ ជាទូទៅ និងរូបមន្តសំខាន់ៗ")
        set_run_font(r_bh, "Khmer OS Muol Light", 11, RGBColor(0, 105, 110), bold=True)

        for th in theories:
            p_th = bcell.add_paragraph()
            p_th.paragraph_format.space_after = Pt(3)
            r_ic = p_th.add_run("➤ ")
            set_run_font(r_ic, "Khmer OS Battambang", 10.5, RGBColor(15, 44, 89), bold=True)
            add_mixed_text(p_th, th)

    sp3 = doc.add_paragraph()
    sp3.paragraph_format.space_before = Pt(4)
    sp3.paragraph_format.space_after = Pt(4)

    # 4. Exercises 2 Columns with multi-page support
    exercises = data.get("exercises", [])
    exercise_pages = []
    curr_page = []
    for i, ex in enumerate(exercises):
        if ex.get("page_break_before") and i > 0 and curr_page:
            exercise_pages.append(curr_page)
            curr_page = [ex]
        else:
            curr_page.append(ex)
    if curr_page:
        exercise_pages.append(curr_page)
    if not exercise_pages and exercises:
        exercise_pages = [exercises]

    ex_counter = 0
    for page_idx, p_exercises in enumerate(exercise_pages):
        if page_idx > 0:
            doc.add_page_break()
            p_cont = doc.add_paragraph()
            p_cont.paragraph_format.space_before = Pt(4)
            p_cont.paragraph_format.space_after = Pt(8)
            r_c = p_cont.add_run(f"▌ {data.get('lesson_title', '')} ៖ លំហាត់អនុវត្តបន្ត (ទំព័រទី {page_idx + 1})")
            set_run_font(r_c, "Khmer OS Muol Light", 11, RGBColor(15, 44, 89), bold=True)

        pairs = [p_exercises[i:i+2] for i in range(0, len(p_exercises), 2)]
        t_ex = doc.add_table(rows=len(pairs), cols=2)
        t_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_ex.autofit = False

        col_w = Cm(8.8)
        for r_idx, pair in enumerate(pairs):
            row = t_ex.rows[r_idx]
            for c_idx in range(2):
                cell = row.cells[c_idx]
                cell.width = col_w
                if c_idx == 0:
                    set_cell_margins(cell, 120, 140, 60, 140)
                    set_cell_border(cell, right={'val': 'single', 'sz': '6', 'color': 'DCD4C6'})
                else:
                    set_cell_margins(cell, 120, 140, 140, 60)
                    set_cell_border(cell)

                if c_idx < len(pair):
                    ex_item = pair[c_idx]
                    ex_counter += 1
                    p_ex = cell.paragraphs[0]
                    p_ex.paragraph_format.space_after = Pt(4)
                    num_str = str(ex_item.get("num", str(ex_counter))).zfill(2)
                    r_d1 = p_ex.add_run(num_str[0])
                    set_run_font(r_d1, "Khmer OS Battambang", 22, RGBColor(115, 125, 135), bold=True)
                    r_d2 = p_ex.add_run(num_str[1] + " ")
                    set_run_font(r_d2, "Khmer OS Battambang", 22, RGBColor(235, 115, 30), bold=True)
                    add_mixed_text(p_ex, ex_item.get("text", ""))

                    # Geometry Diagram preset insertion
                    diag_key = ex_item.get("diagram")
                    if diag_key and diag_key in DIAGRAM_PRESETS:
                        try:
                            png_bytes = render_diagram_png(diag_key)
                            img_buf = io.BytesIO(png_bytes)
                            p_img = cell.add_paragraph()
                            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_img.paragraph_format.space_before = Pt(2)
                            p_img.paragraph_format.space_after = Pt(2)
                            p_img.add_run().add_picture(img_buf, width=Cm(3.8))
                        except Exception as e:
                            print(f"Error embedding diagram: {e}")

                    # Custom drawn image from interactive canvas
                    custom_img = ex_item.get("custom_image")
                    if custom_img and custom_img.startswith("data:image"):
                        try:
                            b64_data = re.sub(r'^data:image/.+;base64,', '', custom_img)
                            img_bytes = base64.b64decode(b64_data)
                            img_buf = io.BytesIO(img_bytes)
                            p_cimg = cell.add_paragraph()
                            p_cimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_cimg.paragraph_format.space_before = Pt(2)
                            p_cimg.paragraph_format.space_after = Pt(2)
                            p_cimg.add_run().add_picture(img_buf, width=Cm(3.8))
                        except Exception as e:
                            print(f"Error embedding custom image: {e}")

                    subs = ex_item.get("subs", [])
                    for sub in subs:
                        if sub.get("text", "").strip():
                            p_sub = cell.add_paragraph()
                            p_sub.paragraph_format.space_after = Pt(2.5)
                            badge_label = sub.get("label", "ក")
                            add_badge(p_sub, badge_label)
                            add_mixed_text(p_sub, sub.get("text", ""))

                    # Teacher Solution insertion
                    sol = ex_item.get("solution")
                    if sol and data.get("show_solutions"):
                        p_sol = cell.add_paragraph()
                        p_sol.paragraph_format.space_before = Pt(3)
                        p_sol.paragraph_format.space_after = Pt(2)
                        r_sol = p_sol.add_run("💡 ដំណោះស្រាយគំរូ ៖ ")
                        set_run_font(r_sol, "Khmer OS Battambang", 9.5, RGBColor(0, 105, 110), bold=True)
                        add_mixed_text(p_sol, sol, default_font="Khmer OS Battambang", font_size=9.5, text_color=RGBColor(0, 105, 110))

    sp4 = doc.add_paragraph()
    sp4.paragraph_format.space_before = Pt(6)
    sp4.paragraph_format.space_after = Pt(4)

    # 5. Motto Banner
    t_motto = doc.add_table(rows=1, cols=1)
    t_motto.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_motto.autofit = False
    mcell = t_motto.rows[0].cells[0]
    mcell.width = Cm(18.0)
    set_cell_shading(mcell, "E6F5F5")
    set_cell_margins(mcell, 80, 80, 100, 100)
    set_cell_border(mcell, top={'val': 'single', 'sz': '4', 'color': '00969B'},
                           bottom={'val': 'single', 'sz': '4', 'color': '00969B'},
                           left={'val': 'single', 'sz': '4', 'color': '00969B'},
                           right={'val': 'single', 'sz': '4', 'color': '00969B'})
    pm = mcell.paragraphs[0]
    pm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = pm.add_run(data.get("motto", "➤ ខំរៀនដើម្បីខ្លួនឯង និងអនាគតដ៏ភ្លឺស្វាង ➤"))
    set_run_font(rm, "Khmer OS Muol Light", 10, RGBColor(15, 44, 89), bold=True)

    # Footer
    footer = sec.footer
    ft = footer.add_table(rows=1, cols=3, width=Cm(18.0))
    ft.alignment = WD_TABLE_ALIGNMENT.CENTER
    ft.autofit = False
    fc0, fc1, fc2 = ft.rows[0].cells
    fc0.width = Cm(6.0)
    fc1.width = Cm(6.0)
    fc2.width = Cm(6.0)
    for c in (fc0, fc1, fc2):
        set_cell_border(c, top={'val': 'single', 'sz': '4', 'color': '00969B'})

    fp0 = fc0.paragraphs[0]
    fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rf0 = fp0.add_run(f"1 | {data.get('lesson_title', 'ចំណងជើង')}")
    set_run_font(rf0, "Khmer OS Battambang", 8.5, RGBColor(15, 44, 89), bold=True)

    fp1 = fc1.paragraphs[0]
    fp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf1 = fp1.add_run(f"★ បង្រៀនដោយ ៖ {data.get('teacher', 'ក្រូត រស្មី')} ★")
    set_run_font(rf1, "Khmer OS Battambang", 8.5, RGBColor(0, 105, 110), bold=True)

    fp2 = fc2.paragraphs[0]
    fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rf2 = fp2.add_run(f"Telegram: {data.get('phone', '070897540')} | 1")
    set_run_font(rf2, "Khmer OS Battambang", 8.5, RGBColor(195, 35, 60), bold=True)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def build_pdf_from_data(data):
    """
    Generates a pixel-perfect XeLaTeX PDF from lesson data.
    """
    lesson_num = data.get("lesson_num", "មេរៀនទី ...")
    lesson_title = data.get("lesson_title", "ចំណងជើងមេរៀន")
    sec_num = data.get("section_num", "១")
    sec_title = data.get("section_title", "ចំណងជើងផ្នែក")
    subsec_title = data.get("subsec_title", "១.១ ចំណងជើងផ្នែករង")
    teacher = data.get("teacher", "ក្រូត រស្មី")
    phone = data.get("phone", "070897540")
    motto = data.get("motto", "➤ ខំរៀនដើម្បីខ្លួនឯង និងអនាគតដ៏ភ្លឺស្វាង ➤")
    theories = data.get("theories", [])
    exercises = data.get("exercises", [])

    th_items = "\n".join([f"    \\item[\\color{{brandNavy}}\\ding{{228}}] \\textbf{{{t}}}" for t in theories if t.strip()])

    temp_dir = tempfile.mkdtemp()

    # Split exercises into pages by page_break_before
    exercise_pages = []
    curr_page = []
    for i, ex in enumerate(exercises):
        if ex.get("page_break_before") and i > 0 and curr_page:
            exercise_pages.append(curr_page)
            curr_page = [ex]
        else:
            curr_page.append(ex)
    if curr_page:
        exercise_pages.append(curr_page)
    if not exercise_pages and exercises:
        exercise_pages = [exercises]

    global_ex_idx = 0

    def render_ex_col(ex_list):
        nonlocal global_ex_idx
        out = []
        for ex in ex_list:
            global_ex_idx += 1
            num_val = str(ex.get("num", str(global_ex_idx))).lstrip("0") or "1"
            text = ex.get("text", "")

            # Diagram: Preset TikZ or Custom Canvas PNG
            diag_key = ex.get("diagram")
            diag_code = ""
            if diag_key and diag_key in DIAGRAM_PRESETS:
                tikz = DIAGRAM_PRESETS[diag_key].get("tikz", "")
                diag_code = f"\n  \\par\\vspace{{2pt}}{{\\centering {tikz} \\par}}\\vspace{{2pt}}"

            custom_img = ex.get("custom_image")
            if custom_img and custom_img.startswith("data:image"):
                try:
                    b64_data = re.sub(r'^data:image/.+;base64,', '', custom_img)
                    img_bytes = base64.b64decode(b64_data)
                    img_filename = f"canvas_diag_{global_ex_idx}.png"
                    with open(os.path.join(temp_dir, img_filename), "wb") as f_img:
                        f_img.write(img_bytes)
                    diag_code = f"\n  \\par\\vspace{{2pt}}{{\\centering \\includegraphics[width=3.8cm]{{{img_filename}}} \\par}}\\vspace{{2pt}}"
                except Exception as e:
                    print(f"Error saving custom canvas image for PDF: {e}")

            subs = ex.get("subs", [])
            sub_items = "\n".join([f"    \\item[\\alphabadge{{{s.get('label', 'ក')}}}] {s.get('text', '')}" for s in subs if s.get('text', '').strip()])
            subs_block = f"""  \\begin{{itemize}}[leftmargin=15pt, itemsep=2pt, topsep=1pt]
{sub_items}
  \\end{{itemize}}""" if sub_items else ""

            # Teacher solution
            sol = ex.get("solution")
            sol_code = ""
            if sol and data.get("show_solutions"):
                sol_code = f"\n  \\par\\vspace{{2pt}}{{\\fontsize{{9}}{{9}}\\selectfont\\color{{koreaDarkTeal}}\\textbf{{\\ding{{72}}\\ ដំណោះស្រាយគំរូ ៖}} {sol}\\par}}"

            out.append(f"""  \\noindent
  \\exhead{{{num_val}}} {text}\\par{diag_code}
{subs_block}{sol_code}
  \\vspace{{5pt}}""")
        return "\n\n".join(out)

    pages_tex_blocks = []
    for p_i, p_exercises in enumerate(exercise_pages):
        half = (len(p_exercises) + 1) // 2
        col1 = p_exercises[:half]
        col2 = p_exercises[half:]
        c1_code = render_ex_col(col1)
        c2_code = render_ex_col(col2)
        
        prefix = ""
        if p_i > 0:
            prefix = f"\\newpage\n\\noindent{{\\color{{brandNavy}}\\textbf{{\\large ▌ \\LessonTitle\\ ៖ លំហាត់អនុវត្តបន្ត (ទំព័រទី {p_i + 1})}}}}\\par\\vspace{{6pt}}\n"
        block = f"""{prefix}\\begin{{multicols}}{{2}}
{c1_code}

  \\columnbreak

{c2_code}
\\end{{multicols}}"""
        pages_tex_blocks.append(block)

    all_exercises_tex = "\n\n\\vspace{4pt}\n\n".join(pages_tex_blocks)

    template_file = os.path.join(PROJECT_DIR, "lesson_template.tex")
    if os.path.exists(template_file):
        with open(template_file, "r", encoding="utf-8") as f:
            base_tex = f.read()
    else:
        raise FileNotFoundError("lesson_template.tex not found")

    # Replace macros
    tex = base_tex
    tex = tex.replace(r"\newcommand{\LessonNum}{មេរៀនទី ...}", f"\\newcommand{{\\LessonNum}}{{{lesson_num}}}")
    tex = tex.replace(r"\newcommand{\LessonTitle}{ចំណងជើងមេរៀន}", f"\\newcommand{{\\LessonTitle}}{{{lesson_title}}}")
    tex = tex.replace(r"\newcommand{\TeacherName}{ក្រូត រស្មី}", f"\\newcommand{{\\TeacherName}}{{{teacher}}}")
    tex = tex.replace(r"\newcommand{\ContactPhone}{070897540}", f"\\newcommand{{\\ContactPhone}}{{{phone}}}")

    # Replace body between \begin{document} and \end{document}
    body_content = f"""\\begin{{document}}

\\koreanheader{{\\LessonNum}}{{\\LessonTitle}}

\\vspace{{3pt}}
\\khsection{{{sec_num}}}{{{sec_title}}}

\\khsubsection{{{subsec_title.split(' ')[0] if ' ' in subsec_title else sec_num + '.1'}}}{{{subsec_title}}}

\\begin{{khboxorange}}
  \\begin{{itemize}}[leftmargin=15pt, itemsep=2pt, topsep=1pt, parsep=0pt]
{th_items}
  \\end{{itemize}}
\\end{{khboxorange}}

\\vspace{{3pt}}

\\setlength{{\\columnsep}}{{18pt}}
\\setlength{{\\columnseprule}}{{0.5pt}}
\\def\\columnseprulecolor{{\\color{{cardBorder}}}}

{all_exercises_tex}

\\vspace{{6pt}}

\\begin{{tcolorbox}}[
  enhanced,
  colback=brandTealBg,
  colframe=brandNavy!30,
  arc=3pt,
  boxrule=0.6pt,
  top=2.5pt, bottom=2.5pt,
  center,
  halign=center
]
  \\fontsize{{9}}{{9}}\\selectfont\\kml\\color{{brandNavy}}
  \\textbf{{\\ding{{228}}\\ {motto} \\ding{{228}}}}
\\end{{tcolorbox}}

\\end{{document}}"""

    # Substitute document body
    doc_start_idx = tex.find(r"\begin{document}")
    if doc_start_idx != -1:
        tex = tex[:doc_start_idx] + body_content
    else:
        tex += "\n" + body_content

    try:
        tex_path = os.path.join(temp_dir, "document.tex")
        with open(tex_path, "w", encoding="utf-8") as f:
            f.write(tex)

        env = os.environ.copy()
        env["TEXINPUTS"] = f".:{PROJECT_DIR}:"
        cmd = ["xelatex", "-interaction=nonstopmode", "document.tex"]
        subprocess.run(cmd, cwd=temp_dir, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

        pdf_path = os.path.join(temp_dir, "document.pdf")
        if os.path.exists(pdf_path):
            with open(pdf_path, "rb") as pf:
                return pf.read()
        else:
            raise RuntimeError("XeLaTeX compilation failed to produce PDF")
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

class LiveEditorHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def do_GET(self):
        if self.path == '/api/diagrams':
            res_bytes = json.dumps(DIAGRAM_PRESETS, ensure_ascii=False).encode('utf-8')
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(res_bytes)))
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(res_bytes)
            return
        return super().do_GET()

    def do_POST(self):
        if self.path == '/api/convert':
            content_len = int(self.headers.get('Content-Length', 0))
            post_body = self.rfile.read(content_len)
            try:
                req_data = json.loads(post_body.decode('utf-8'))
                latex = req_data.get('latex', '').strip()
                mathml = latex2mathml.converter.convert(latex)
                omml = latex_to_omml_xml(latex)
                res = {
                    "success": True,
                    "latex": latex,
                    "mathml": mathml,
                    "omml": omml
                }
            except Exception as e:
                res = {
                    "success": False,
                    "error": str(e)
                }
            res_bytes = json.dumps(res, ensure_ascii=False).encode('utf-8')
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(res_bytes)))
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(res_bytes)
            return

        elif self.path == '/api/export-docx':
            content_len = int(self.headers.get('Content-Length', 0))
            post_body = self.rfile.read(content_len)
            try:
                data = json.loads(post_body.decode('utf-8'))
                docx_bytes = build_docx_from_data(data)
                self.send_response(200)
                self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                self.send_header('Content-Disposition', 'attachment; filename="math_lesson.docx"')
                self.send_header('Content-Length', str(len(docx_bytes)))
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(docx_bytes)
            except Exception as e:
                err_msg = json.dumps({"error": str(e)}).encode('utf-8')
                self.send_response(500)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(err_msg)
            return

        elif self.path == '/api/export-pdf':
            content_len = int(self.headers.get('Content-Length', 0))
            post_body = self.rfile.read(content_len)
            try:
                data = json.loads(post_body.decode('utf-8'))
                pdf_bytes = build_pdf_from_data(data)
                self.send_response(200)
                self.send_header('Content-Type', 'application/pdf')
                self.send_header('Content-Disposition', 'attachment; filename="math_lesson.pdf"')
                self.send_header('Content-Length', str(len(pdf_bytes)))
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(pdf_bytes)
            except Exception as e:
                err_msg = json.dumps({"error": str(e)}).encode('utf-8')
                self.send_response(500)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(err_msg)
            return

        self.send_error(404, "Endpoint not found")

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

if __name__ == '__main__':
    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("", PORT), LiveEditorHandler) as httpd:
        print(f"Server started at http://localhost:{PORT}")
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nShutting down server.")
